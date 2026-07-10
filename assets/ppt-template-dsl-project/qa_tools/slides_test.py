#!/usr/bin/env python3
"""Perform a structural slide-boundary check for shapes in a PPTX."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_path")
    parser.add_argument("--tolerance-emu", type=int, default=12700, help="Allowed boundary tolerance; default is one point")
    parser.add_argument("--all-shapes", action="store_true", help="Also flag decorative and image-filled shapes outside the canvas")
    parser.add_argument("--report-json", help="Write geometry and text-collision diagnostics as JSON")
    parser.add_argument("--fail-on-text-collision", action="store_true", help="Treat text-collision candidates as failures instead of review findings")
    args = parser.parse_args()

    try:
        from pptx import Presentation
    except ImportError as error:
        raise RuntimeError("python-pptx is required; run: python -m pip install -r requirements-qa.txt") from error

    presentation = Presentation(Path(args.input_path).expanduser().resolve())
    failures: list[str] = []
    geometry_failures: list[dict[str, object]] = []
    text_collisions: list[dict[str, object]] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_shapes: list[tuple[int, object, tuple[int, int, int, int]]] = []
        for shape_number, shape in enumerate(slide.shapes, start=1):
            has_visible_text = bool(getattr(shape, "has_text_frame", False) and shape.text.strip())
            if not args.all_shapes and not has_visible_text:
                continue
            left = int(shape.left)
            top = int(shape.top)
            right = left + int(shape.width)
            bottom = top + int(shape.height)
            tolerance = args.tolerance_emu
            if left < -tolerance or top < -tolerance or right > int(presentation.slide_width) + tolerance or bottom > int(presentation.slide_height) + tolerance:
                failures.append(f"slide {slide_number} shape {shape_number}: ({left}, {top}, {right}, {bottom})")
                geometry_failures.append({"slide": slide_number, "shape": shape_number, "bbox": [left, top, right, bottom]})
            if has_visible_text:
                text_shapes.append((shape_number, shape, (left, top, right, bottom)))
        for index, (shape_number, shape, box) in enumerate(text_shapes):
            for other_number, other_shape, other_box in text_shapes[index + 1:]:
                overlap = intersection(box, other_box)
                if overlap <= 0:
                    continue
                smaller = min(area(box), area(other_box))
                if smaller <= 0 or overlap / smaller < 0.35:
                    continue
                collision = {
                    "slide": slide_number,
                    "shape": shape_number,
                    "otherShape": other_number,
                    "overlapRatio": round(overlap / smaller, 4),
                    "text": shape.text.strip()[:80],
                    "otherText": other_shape.text.strip()[:80],
                }
                text_collisions.append(collision)
                if args.fail_on_text_collision:
                    failures.append(f"slide {slide_number} text shapes {shape_number} and {other_number}: overlap ratio {collision['overlapRatio']}")
    if args.report_json:
        report_path = Path(args.report_json).expanduser().resolve()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps({"passed": not failures, "geometryFailures": geometry_failures, "textCollisions": text_collisions, "textCollisionCandidates": len(text_collisions)}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if failures:
        print("ERROR: Slide geometry or text collision failures:\n" + "\n".join(failures), file=sys.stderr)
        sys.exit(1)
    print(f"Test passed. No structural geometry overflow detected; text collision candidates: {len(text_collisions)}.")


def area(box: tuple[int, int, int, int]) -> int:
    return max(0, box[2] - box[0]) * max(0, box[3] - box[1])


def intersection(left: tuple[int, int, int, int], right: tuple[int, int, int, int]) -> int:
    return max(0, min(left[2], right[2]) - max(left[0], right[0])) * max(0, min(left[3], right[3]) - max(left[1], right[1]))


if __name__ == "__main__":
    main()
